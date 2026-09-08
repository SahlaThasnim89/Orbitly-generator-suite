"""Orbitly SalesOS — Product Specifications (Growth Tier) as DOCX.

Paragraphs, version table, module spec tables, integration matrix, SLA table,
and an embedded matplotlib graph (latency vs call volume) for parser tests.
"""

import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
INDIGO = RGBColor(0x1B, 0x2A, 0x4A)
TEAL = RGBColor(0x0D, 0x94, 0x88)
MUTED = RGBColor(0x64, 0x74, 0x8B)


def make_chart():
    fig, ax = plt.subplots(figsize=(6.8, 3.2), dpi=160)
    vols = [100, 500, 1000, 5000, 10000, 25000]
    p50 = [0.9, 1.0, 1.1, 1.4, 1.8, 2.4]
    p99 = [2.1, 2.3, 2.6, 3.2, 4.1, 5.2]
    ax.plot(vols, p50, marker="o", linewidth=2.2, color="#0D9488", label="p50")
    ax.plot(vols, p99, marker="s", linewidth=2.0, color="#1B2A4A", linestyle="--", label="p99")
    ax.set_xscale("log")
    ax.set_xlabel("Calls processed per day (log scale)", fontsize=10)
    ax.set_ylabel("Pipeline latency (minutes)", fontsize=10)
    ax.set_title("Ingestion latency by daily call volume (SLO: p99 < 6 min)", fontsize=11)
    ax.legend(frameon=False, fontsize=10)
    ax.grid(axis="y", color="#E2E8F0", linewidth=0.6)
    ax.set_axisbelow(True)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    p = os.path.join(HERE, "img_ingest_latency.png")
    fig.savefig(p, facecolor="white")
    plt.close(fig)
    return p


CHART = make_chart()

doc = Document()
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.8)

def h1(text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.color.rgb = INDIGO
        r.font.name = "Calibri"

def h2(text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        r.font.color.rgb = INDIGO
        r.font.name = "Calibri"

def para(text, size=11, italic=False, color=None, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.italic = italic
    r.font.bold = bold
    if color:
        r.font.color.rgb = color
    return p

def table(rows, widths=None, header=True):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = str(val)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
                    if ri == 0 and header:
                        r.font.bold = True
    if widths:
        for ci, w in enumerate(widths):
            for row in t.rows:
                row.cells[ci].width = Inches(w)
    doc.add_paragraph()
    return t

# ── Title ────────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Orbitly SalesOS — Product Specifications")
r.font.size = Pt(26)
r.font.bold = True
r.font.color.rgb = INDIGO
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Growth Tier · Platform release v2.6 · August 2026 · Document ID: SPEC-2026-08-G")
r.font.size = Pt(11)
r.font.color.rgb = MUTED

# ── 1. Overview ─────────────────────────────────────────────────────────────
h1("1. Overview")
para("Orbitly SalesOS is a revenue-intelligence platform for B2B sales organizations. It "
     "ingests every sales call (Teams, Zoom, or telephony), transcribes it, and builds a "
     "structured deal graph that connects conversations to CRM records, email threads, and "
     "billing data. Three engines operate on that graph: Conversation Intelligence (call "
     "scoring, objection and competitor extraction), Forecast Intelligence (weekly deal "
     "scoring with confidence and reasons), and Coaching & Ramp (playbook-aligned manager "
     "reviews and onboarding paths). The product is delivered as a native Salesforce "
     "application; reps never leave Salesforce.")
para("This document specifies the Growth tier — the recommended tier for sales organizations "
     "that need forecast scoring and contracted adoption guarantees. Metrics quoted in "
     "customer-facing materials (85%+ forecast accuracy, 3-month ramp, +12% win rate) come "
     "from the FY2026 customer cohort and are detailed in the Orbitly case study pack.", size=10, italic=True, color=MUTED)

# ── 2. Release history ──────────────────────────────────────────────────────
h1("2. Release history")
table([
    ["Release", "Date", "Headline changes"],
    ["v2.6", "Aug 2026", "SAP billing connector GA; forecast reasons v2 (per-deal factor list); Teams ingestion at 2× speed"],
    ["v2.5", "May 2026", "EU data region (Enterprise); rep-facing coaching digest; 14 new playbook scoring dimensions"],
    ["v2.4", "Feb 2026", "Adoption guarantee (Growth); pilot-gate reporting; Salesforce embedded UI refresh"],
    ["v2.3", "Nov 2025", "Zoom telephony connector; objection taxonomy v3; bulk deal rescoring API"],
], widths=[1.0, 1.1, 4.7])

# ── 3. Module specs ─────────────────────────────────────────────────────────
h1("3. Module specifications")

h2("3.1 Conversation Intelligence")
table([
    ["Specification", "Value"],
    ["Transcription languages", "40+ languages; English, Hindi, German, French, Spanish most deployed"],
    ["Transcription word error rate", "< 8% on sales-domain audio (internal eval set, FY2026)"],
    ["Objection taxonomy", "22 categories (price, timing, competitor, security, authority, …) with per-objection confidence"],
    ["Competitor detection", "Automatic; custom competitor watchlists supported per workspace"],
    ["Playbook scoring", "Up to 40 scoring dimensions per call; custom playbooks per team"],
    ["Manager review digests", "Weekly; 2-minute highlight reels ranked by coaching value"],
    ["CRM auto-capture", "Call summary, next steps, objections, competitor mentions → Salesforce fields"],
], widths=[2.6, 4.2])

h2("3.2 Forecast Intelligence")
table([
    ["Specification", "Value"],
    ["Scoring cadence", "Nightly per-deal rescoring; weekly forecast snapshot (Monday 06:00 local)"],
    ["Reported accuracy", "60% baseline → 85%+ within two quarters (customer cohort median)"],
    ["Per-deal output", "Confidence 0–100, reason factors, slippage risk flags, pull-through estimate"],
    ["Data inputs", "Calls, email threads, CRM stage history, SAP billing/fulfillment signals"],
    ["Scenario views", "Commit / best case / pipeline rollups with variance-vs-last-week"],
    ["Alerting", "Slippage and stalling alerts to owner + manager (email, Salesforce, Teams)"],
], widths=[2.6, 4.2])

h2("3.3 Coaching & Ramp")
table([
    ["Specification", "Value"],
    ["Ramp compression", "Median 6 months → 3 months to full quota (cohort median)"],
    ["Onboarding paths", "Role-based paths with call milestones and auto-graded practice pitches"],
    ["Review workflow", "Manager review target < 15 minutes per rep per week"],
    ["Rep-facing mode", "Self-coaching digest first; manager visibility configurable per team"],
], widths=[2.6, 4.2])

# ── 4. Integrations ─────────────────────────────────────────────────────────
h1("4. Integrations")
table([
    ["System", "Mode", "Notes"],
    ["Salesforce", "Native app, bi-directional", "Embedded UI; no separate rep login. objects: Opportunity, Task, Call2, custom"],
    ["SAP (SD/FI)", "Certified connector (v2.6 GA)", "Billing + fulfillment reconciliation into forecast factors"],
    ["Microsoft Teams", "Built-in ingestion", "Auto-join or bot-join; recordings processed at 2× real time"],
    ["Zoom", "Built-in ingestion", "Cloud recording connector; telephony bridge supported"],
    ["Gmail / Outlook 365", "Email graph ingestion", "Thread-level sentiment and commitment extraction"],
    ["REST API + Webhooks", "v2, JSON", "Deal scores, call metadata, objection events; 600 req/min per workspace"],
], widths=[1.7, 1.9, 3.2])

# ── 5. Limits & SLA ─────────────────────────────────────────────────────────
h1("5. Limits, retention, and SLA")
table([
    ["Item", "Team", "Growth", "Enterprise"],
    ["Seats (min–max)", "10–250", "25–1,000", "Unlimited"],
    ["Call audio retention", "12 months", "24 months", "36 months, customer-managed keys"],
    ["API rate limit", "200 req/min", "600 req/min", "Custom"],
    ["Data residency", "US", "In-region add-on (India/EU)", "Included (India/EU/US)"],
    ["Uptime SLA", "—", "99.9%", "99.95% + priority support"],
    ["Adoption guarantee", "—", "Contracted", "Contracted"],
    ["Price / seat / year", "$480", "$600", "Custom"],
], widths=[2.1, 1.5, 1.7, 1.5])

# ── 6. Ingestion performance graph ─────────────────────────────────────────
h1("6. Ingestion performance")
para("Pipeline latency from call end to searchable transcript and scored deal, by daily "
     "ingestion volume. The p99 SLO of six minutes holds to roughly 25,000 calls/day on a "
     "standard Growth deployment; above that, sharded ingestion is provisioned. Figure 1 "
     "below shows the measured cohort medians for release v2.6.")
doc.add_picture(CHART, width=Inches(6.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
cap = para("Figure 1 — Ingestion latency (minutes) by daily call volume, v2.6, FY2026 cohort.",
           size=9, italic=True, color=MUTED)
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

# ── 7. Security ─────────────────────────────────────────────────────────────
h1("7. Security and compliance")
para("Orbitly maintains SOC 2 Type II and ISO 27001 certifications. Customer call data is "
     "never used to train any model. Data residency for India and the EU is included on the "
     "Enterprise tier and available as an in-region processing add-on on Growth. A signed "
     "DPA with EU Standard Contractual Clauses is provided in the security pack, alongside "
     "penetration-test summaries and the current sub-processor list.")

out = os.path.join(HERE, "orbitly_product_specs.docx")
doc.save(out)
print(f"WROTE {out}")
